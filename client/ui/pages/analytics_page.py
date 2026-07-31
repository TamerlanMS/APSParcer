# -*- coding: utf-8 -*-
"""
Фаза 4 — Аналитика.
Пять вкладок: Обзор | Бренды | ИИ-подбор | Цены | Аномалии
+ Экспорт в Excel (клиентская генерация)
+ Настраиваемый период (7/30/90 дн. или ввод вручную)
"""
import threading
import datetime
import customtkinter as ctk
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from assets.theme import *
from locales.strings import t

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False


# ── Цвета карточек KPI ────────────────────────────────────────────────────────
CARD_COLORS = [
    ("#2C3E50", "#ECF0F1"),
    ("#1A6B3A", "#E9F7EF"),
    ("#145A8C", "#EBF5FB"),
    ("#7B241C", "#FDEDEC"),
]

PERIOD_OPTIONS = [
    ("7 дн.",  7),
    ("30 дн.", 30),
    ("90 дн.", 90),
]

_PRICE_FIELDS = {
    "kaznisa": "КазНИИСА",
    "rrts":    "РРЦС",
    "mrc":     "МРЦ",
    "opt":     "Оптовая",
}

_STATUS_NAMES = {
    "exact_match":  "Точное совпадение",
    "code_exact":   "Точное совпадение (артикул)",
    "fuzzy":        "Нечёткое совпадение",
    "ai_match":     "ИИ-подбор",
    "not_found":    "Не найдено",
    "correction":   "Из истории исправлений",
    "heading":      "Заголовок секции",
    "unknown":      "Неизвестно",
}


class _KpiCard(ctk.CTkFrame):
    """Одна сводная карточка: иконка + число + подпись."""

    def __init__(self, parent, icon: str, label: str, accent: str, **kw):
        super().__init__(parent, fg_color=BG_CARD, corner_radius=RADIUS_MD,
                         border_width=2, border_color=accent, **kw)
        top = ctk.CTkFrame(self, fg_color=accent, corner_radius=RADIUS_SM,
                           width=42, height=42)
        top.pack(anchor="w", padx=14, pady=(14, 6))
        top.pack_propagate(False)
        ctk.CTkLabel(top, text=icon, font=("Segoe UI Emoji", 18),
                     text_color="white").place(relx=0.5, rely=0.5, anchor="center")
        self._value_lbl = ctk.CTkLabel(self, text="—",
                                        font=("Calibri", 22, "bold"),
                                        text_color=accent)
        self._value_lbl.pack(anchor="w", padx=14)
        ctk.CTkLabel(self, text=label, font=FONT_SMALL,
                     text_color=TEXT_SECONDARY).pack(anchor="w", padx=14, pady=(2, 12))

    def set(self, value):
        self._value_lbl.configure(text=str(value))


class AnalyticsPage(ctk.CTkFrame):
    """Страница аналитики (Фаза 4) — пять вкладок."""

    # Tab names (constants to avoid typos)
    _T_OVERVIEW  = "📊  Обзор"
    _T_BRANDS    = "🏷  Бренды"
    _T_AI        = "🤖  ИИ-подбор"
    _T_PRICES    = "🔍  Цены"
    _T_ANOMALIES = "⚠  Аномалии"

    def __init__(self, parent, api, app):
        super().__init__(parent, fg_color=BG_MAIN, corner_radius=0)
        self.api = api
        self.app = app
        self._period = 30
        self._loading = False
        # Last fetched data (for export)
        self._last_summary: dict = {}
        self._last_kpi:     dict = {}
        self._last_brands:  dict = {}
        self._last_ai_eff:  dict = {}
        self._last_anomalies: dict = {}
        self._build()

    # ── Построение интерфейса ─────────────────────────────────────────────────

    def _build(self):
        self.grid_rowconfigure(2, weight=1)
        self.grid_columnconfigure(0, weight=1)

        # ── Шапка ─────────────────────────────────────────────────────────────
        hdr = ctk.CTkFrame(self, fg_color=NAVY_DARK, corner_radius=0)
        hdr.grid(row=0, column=0, sticky="ew")
        hdr.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(
            hdr, text="📈  Аналитика",
            font=FONT_HEADING, text_color="white",
        ).grid(row=0, column=0, padx=20, pady=14, sticky="w")

        # Правый блок: период + кастом + обновить + экспорт
        right_frame = ctk.CTkFrame(hdr, fg_color="transparent")
        right_frame.grid(row=0, column=2, padx=20, pady=10, sticky="e")

        ctk.CTkLabel(right_frame, text="Период:", font=FONT_SMALL,
                     text_color=TEXT_NAV).pack(side="left", padx=(0, 6))

        period_labels = [p[0] for p in PERIOD_OPTIONS]
        self._period_var = ctk.StringVar(value=period_labels[1])
        self._period_seg = ctk.CTkSegmentedButton(
            right_frame, values=period_labels, variable=self._period_var,
            font=FONT_SMALL, selected_color=BLUE_MID,
            selected_hover_color=NAVY_DARK, unselected_color="#1E3D6B",
            unselected_hover_color=NAVY_DARK, text_color="white",
            text_color_disabled="#8899AA", height=28,
            command=self._on_period_change,
        )
        self._period_seg.pack(side="left")

        # Поле кастомного периода
        ctk.CTkLabel(right_frame, text="или", font=FONT_SMALL,
                     text_color=TEXT_NAV).pack(side="left", padx=(10, 4))
        self._custom_days_var = ctk.StringVar(value="")
        custom_entry = ctk.CTkEntry(
            right_frame, textvariable=self._custom_days_var,
            width=52, height=28, font=FONT_SMALL,
            placeholder_text="дней",
        )
        custom_entry.pack(side="left")
        custom_entry.bind("<Return>", lambda _: self._apply_custom_period())
        ctk.CTkButton(
            right_frame, text="→", font=FONT_SMALL,
            fg_color="#1E3D6B", hover_color=BLUE_MID,
            text_color="white", width=28, height=28,
            corner_radius=RADIUS_SM, command=self._apply_custom_period,
        ).pack(side="left", padx=(2, 0))

        ctk.CTkButton(
            right_frame, text="↻", font=("Calibri", 15, "bold"),
            fg_color="#1E3D6B", hover_color=BLUE_MID,
            text_color="white", width=32, height=28,
            corner_radius=RADIUS_SM, command=self.load,
        ).pack(side="left", padx=(10, 0))

        ctk.CTkButton(
            right_frame, text="⬇  Excel", font=FONT_SMALL,
            fg_color="#1A6B3A", hover_color="#145A32",
            text_color="white", height=28, corner_radius=RADIUS_SM,
            command=self._export_to_excel,
        ).pack(side="left", padx=(8, 0))

        # ── Статус-строка ──────────────────────────────────────────────────────
        self._status_lbl = ctk.CTkLabel(
            self, text="", font=FONT_SMALL, text_color=TEXT_SECONDARY)
        self._status_lbl.grid(row=1, column=0, pady=(4, 0))

        # ── Общий стиль Treeview ──────────────────────────────────────────────
        style = ttk.Style()
        style.theme_use("clam")
        style.configure("Analytics.Treeview",
                        background=BG_CARD, foreground="#2C3E50",
                        rowheight=34, fieldbackground=BG_CARD,
                        bordercolor="#DDE2E8", relief="flat",
                        font=("Calibri", 12))
        style.configure("Analytics.Treeview.Heading",
                        background=NAVY_DARK, foreground="white",
                        font=("Calibri", 11, "bold"), relief="flat")
        style.map("Analytics.Treeview",
                  background=[("selected", BLUE_MID)],
                  foreground=[("selected", "white")])
        style.map("Analytics.Treeview.Heading", relief=[("active", "flat")])

        # ── Вкладки ───────────────────────────────────────────────────────────
        self._tabs = ctk.CTkTabview(
            self, fg_color=BG_MAIN,
            segmented_button_fg_color=NAVY_DARK,
            segmented_button_selected_color=BLUE_MID,
            segmented_button_selected_hover_color="#1E5CA8",
            segmented_button_unselected_color=NAVY_DARK,
            segmented_button_unselected_hover_color="#1E3D6B",
            text_color="white",
            text_color_disabled="#8899AA",
            corner_radius=0,
        )
        self._tabs.grid(row=2, column=0, sticky="nsew")

        for name in (self._T_OVERVIEW, self._T_BRANDS, self._T_AI,
                     self._T_PRICES, self._T_ANOMALIES):
            self._tabs.add(name)
        self._tabs.set(self._T_OVERVIEW)

        self._build_overview_tab(self._tabs.tab(self._T_OVERVIEW))
        self._build_brands_tab(self._tabs.tab(self._T_BRANDS))
        self._build_ai_tab(self._tabs.tab(self._T_AI))
        self._build_prices_tab(self._tabs.tab(self._T_PRICES))
        self._build_anomalies_tab(self._tabs.tab(self._T_ANOMALIES))

    # ── Вкладка 1: Обзор ──────────────────────────────────────────────────────

    def _build_overview_tab(self, tab):
        tab.grid_columnconfigure(0, weight=1)
        tab.grid_rowconfigure(3, weight=1)

        cards_frame = ctk.CTkFrame(tab, fg_color="transparent")
        cards_frame.grid(row=0, column=0, sticky="ew", padx=PAD_MD, pady=(12, 0))
        for i in range(4):
            cards_frame.grid_columnconfigure(i, weight=1, uniform="card")

        icons   = ["📋", "📦", "👤", "⚠️"]
        labels  = ["КП за период", "Позиций обработано", "Активных менеджеров", "Ошибок"]
        accents = [c[0] for c in CARD_COLORS]

        self._cards: list[_KpiCard] = []
        for i, (icon, label, accent) in enumerate(zip(icons, labels, accents)):
            card = _KpiCard(cards_frame, icon=icon, label=label, accent=accent)
            card.grid(row=0, column=i, sticky="nsew", padx=6, pady=6)
            self._cards.append(card)

        self._avg_lbl = ctk.CTkLabel(
            tab, text="", font=FONT_SMALL, text_color=TEXT_SECONDARY)
        self._avg_lbl.grid(row=1, column=0, sticky="w", padx=PAD_MD + 6, pady=(4, 0))

        tbl_hdr = ctk.CTkFrame(tab, fg_color="transparent")
        tbl_hdr.grid(row=2, column=0, sticky="ew", padx=PAD_MD, pady=(16, 0))
        ctk.CTkLabel(tbl_hdr, text="Анализ менеджеров",
                     font=FONT_HEADING, text_color=NAVY_DARK).pack(side="left")

        tbl_wrap = ctk.CTkFrame(tab, fg_color=BG_CARD,
                                corner_radius=RADIUS_MD, border_width=1,
                                border_color="#DDE2E8")
        tbl_wrap.grid(row=3, column=0, sticky="nsew", padx=PAD_MD, pady=(8, PAD_MD))
        tbl_wrap.grid_columnconfigure(0, weight=1)
        tbl_wrap.grid_rowconfigure(0, weight=1)

        cols = ("full_name", "total_kp", "total_items", "avg_items", "last_activity")
        self._tree = ttk.Treeview(tbl_wrap, columns=cols, show="headings",
                                   style="Analytics.Treeview", selectmode="browse",
                                   height=12)
        col_cfg = [
            ("full_name",     "ФИО менеджера",        220, "w"),
            ("total_kp",      "КП за период",          120, "center"),
            ("total_items",   "Позиций",               120, "center"),
            ("avg_items",     "Ср. поз./КП",           110, "center"),
            ("last_activity", "Последняя активность",  160, "center"),
        ]
        for col, heading, width, anchor in col_cfg:
            self._tree.heading(col, text=heading,
                               command=lambda c=col: self._sort_tree(c, False))
            self._tree.column(col, width=width, anchor=anchor, minwidth=80)
        self._tree.tag_configure("odd",  background="#F8FAFB")
        self._tree.tag_configure("even", background=BG_CARD)

        vsb = ttk.Scrollbar(tbl_wrap, orient="vertical", command=self._tree.yview)
        self._tree.configure(yscrollcommand=vsb.set)
        self._tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")

        self._empty_lbl = ctk.CTkLabel(
            tab, text="Нет данных за выбранный период",
            font=FONT_NORMAL, text_color=TEXT_SECONDARY)

    # ── Вкладка 2: Бренды ─────────────────────────────────────────────────────

    def _build_brands_tab(self, tab):
        tab.grid_columnconfigure(0, weight=1)
        tab.grid_rowconfigure(1, weight=1)

        ctk.CTkLabel(
            tab, font=FONT_SMALL, text_color=TEXT_SECONDARY,
            text="Топ брендов по числу ручных исправлений менеджеров за период",
        ).grid(row=0, column=0, sticky="w", padx=PAD_MD, pady=(12, 0))

        brd_wrap = ctk.CTkFrame(tab, fg_color=BG_CARD,
                                corner_radius=RADIUS_MD, border_width=1,
                                border_color="#DDE2E8")
        brd_wrap.grid(row=1, column=0, sticky="nsew", padx=PAD_MD, pady=(8, 0))
        brd_wrap.grid_columnconfigure(0, weight=1)
        brd_wrap.grid_rowconfigure(0, weight=1)

        cols = ("brand", "corrections", "pct", "unique_users")
        self._brand_tree = ttk.Treeview(brd_wrap, columns=cols, show="headings",
                                         style="Analytics.Treeview", selectmode="browse",
                                         height=14)
        for col, heading, width, anchor in [
            ("brand",        "Бренд",           280, "w"),
            ("corrections",  "Исправлений",      140, "center"),
            ("pct",          "%",                90,  "center"),
            ("unique_users", "Пользователей",    140, "center"),
        ]:
            self._brand_tree.heading(col, text=heading)
            self._brand_tree.column(col, width=width, anchor=anchor, minwidth=60)
        self._brand_tree.tag_configure("odd",  background="#F8FAFB")
        self._brand_tree.tag_configure("even", background=BG_CARD)

        vsb_b = ttk.Scrollbar(brd_wrap, orient="vertical",
                               command=self._brand_tree.yview)
        self._brand_tree.configure(yscrollcommand=vsb_b.set)
        self._brand_tree.grid(row=0, column=0, sticky="nsew")
        vsb_b.grid(row=0, column=1, sticky="ns")

        self._brand_status_lbl = ctk.CTkLabel(
            tab, text="", font=FONT_SMALL, text_color=TEXT_SECONDARY,
            wraplength=900)
        self._brand_status_lbl.grid(row=2, column=0, sticky="w",
                                    padx=PAD_MD + 6, pady=(6, PAD_MD))

    # ── Вкладка 3: ИИ-подбор ─────────────────────────────────────────────────

    def _build_ai_tab(self, tab):
        tab.grid_columnconfigure(0, weight=1)
        tab.grid_rowconfigure(3, weight=1)

        ctk.CTkLabel(
            tab, font=FONT_SMALL, text_color=TEXT_SECONDARY,
            text="Анализ ручных вмешательств менеджеров и точности AI-матчинга",
        ).grid(row=0, column=0, sticky="w", padx=PAD_MD, pady=(12, 0))

        ai_cards = ctk.CTkFrame(tab, fg_color="transparent")
        ai_cards.grid(row=1, column=0, sticky="ew", padx=PAD_MD, pady=(8, 0))
        for i in range(4):
            ai_cards.grid_columnconfigure(i, weight=1, uniform="aicard")

        self._ai_val_lbls = []
        for i, (icon, lbl_text, accent) in enumerate([
            ("📦", "Позиций обработано",    "#145A8C"),
            ("✏️",  "Ручных исправлений",   "#1A6B3A"),
            ("📊", "Частота исправлений",   "#7B241C"),
            ("🔄", "Проиндексировано в AI", "#2C3E50"),
        ]):
            card = ctk.CTkFrame(ai_cards, fg_color=BG_CARD, corner_radius=RADIUS_SM,
                                border_width=1, border_color=accent)
            card.grid(row=0, column=i, sticky="nsew", padx=4)
            top_f = ctk.CTkFrame(card, fg_color=accent, corner_radius=RADIUS_SM,
                                 width=34, height=34)
            top_f.pack(anchor="w", padx=10, pady=(10, 4))
            top_f.pack_propagate(False)
            ctk.CTkLabel(top_f, text=icon, font=("Segoe UI Emoji", 14),
                         text_color="white").place(relx=0.5, rely=0.5, anchor="center")
            val = ctk.CTkLabel(card, text="—", font=("Calibri", 18, "bold"),
                               text_color=accent)
            val.pack(anchor="w", padx=10)
            ctk.CTkLabel(card, text=lbl_text, font=FONT_SMALL,
                         text_color=TEXT_SECONDARY, wraplength=130).pack(
                         anchor="w", padx=10, pady=(2, 10))
            self._ai_val_lbls.append(val)

        ai_tbl_hdr = ctk.CTkFrame(tab, fg_color="transparent")
        ai_tbl_hdr.grid(row=2, column=0, sticky="ew", padx=PAD_MD, pady=(16, 0))
        ctk.CTkLabel(ai_tbl_hdr, text="Разбивка по исходному статусу",
                     font=FONT_HEADING, text_color=NAVY_DARK).pack(side="left")

        ai_tbl_wrap = ctk.CTkFrame(tab, fg_color=BG_CARD,
                                   corner_radius=RADIUS_MD, border_width=1,
                                   border_color="#DDE2E8")
        ai_tbl_wrap.grid(row=3, column=0, sticky="nsew", padx=PAD_MD, pady=(8, 0))
        ai_tbl_wrap.grid_columnconfigure(0, weight=1)
        ai_tbl_wrap.grid_rowconfigure(0, weight=1)

        ai_cols = ("status", "count", "pct")
        self._ai_tree = ttk.Treeview(ai_tbl_wrap, columns=ai_cols, show="headings",
                                      style="Analytics.Treeview", selectmode="browse",
                                      height=7)
        for col, heading, width, anchor in [
            ("status", "Исходный статус (до вмешательства менеджера)", 380, "w"),
            ("count",  "Исправлений",                                   140, "center"),
            ("pct",    "%",                                              90,  "center"),
        ]:
            self._ai_tree.heading(col, text=heading)
            self._ai_tree.column(col, width=width, anchor=anchor, minwidth=60)
        self._ai_tree.tag_configure("odd",  background="#F8FAFB")
        self._ai_tree.tag_configure("even", background=BG_CARD)
        self._ai_tree.grid(row=0, column=0, sticky="nsew")

        ctk.CTkLabel(
            tab,
            text=("ℹ️  «Не найдено» — ИИ не смог подобрать, менеджер нашёл сам. "
                  "«ИИ-подбор» — ИИ предложил вариант, но менеджер его заменил."),
            font=FONT_SMALL, text_color=TEXT_SECONDARY,
            wraplength=820, justify="left",
        ).grid(row=4, column=0, sticky="w", padx=PAD_MD + 6, pady=(6, PAD_MD))

    # ── Вкладка 4: Цены ───────────────────────────────────────────────────────

    def _build_prices_tab(self, tab):
        tab.grid_columnconfigure(0, weight=1)
        tab.grid_rowconfigure(2, weight=1)

        # Поиск
        search_frame = ctk.CTkFrame(tab, fg_color=BG_CARD,
                                    corner_radius=RADIUS_MD, border_width=1,
                                    border_color="#DDE2E8")
        search_frame.grid(row=0, column=0, sticky="ew", padx=PAD_MD, pady=(12, 0))

        ctk.CTkLabel(search_frame, text="Артикул:", font=FONT_SMALL,
                     text_color=TEXT_SECONDARY).grid(row=0, column=0, padx=(14, 6), pady=10)

        self._price_article_var = ctk.StringVar()
        article_entry = ctk.CTkEntry(
            search_frame, textvariable=self._price_article_var,
            width=240, height=32, font=FONT_NORMAL,
            placeholder_text="напр. ABB-1234 или 2CDS271001R0064",
        )
        article_entry.grid(row=0, column=1, padx=6, pady=10)
        article_entry.bind("<Return>", lambda _: self._search_price_history())
        self._bind_paste(article_entry)

        ctk.CTkLabel(search_frame, text="Сегмент:", font=FONT_SMALL,
                     text_color=TEXT_SECONDARY).grid(row=0, column=2, padx=(12, 6))

        self._price_segment_var = ctk.StringVar(value="Все")
        ctk.CTkOptionMenu(
            search_frame, values=["Все", "ss", "os", "sil"],
            variable=self._price_segment_var, font=FONT_SMALL,
            width=90, height=32,
        ).grid(row=0, column=3, padx=6)

        ctk.CTkButton(
            search_frame, text="Найти", font=FONT_SMALL,
            fg_color=BLUE_MID, hover_color=NAVY_DARK,
            text_color="white", height=32, width=80,
            corner_radius=RADIUS_SM, command=self._search_price_history,
        ).grid(row=0, column=4, padx=(6, 14))

        self._price_status_lbl = ctk.CTkLabel(
            tab, text="Введите артикул и нажмите «Найти»",
            font=FONT_SMALL, text_color=TEXT_SECONDARY)
        self._price_status_lbl.grid(row=1, column=0, sticky="w",
                                    padx=PAD_MD + 6, pady=(6, 0))

        # Таблица истории цен
        ph_wrap = ctk.CTkFrame(tab, fg_color=BG_CARD,
                               corner_radius=RADIUS_MD, border_width=1,
                               border_color="#DDE2E8")
        ph_wrap.grid(row=2, column=0, sticky="nsew", padx=PAD_MD, pady=(6, PAD_MD))
        ph_wrap.grid_columnconfigure(0, weight=1)
        ph_wrap.grid_rowconfigure(0, weight=1)

        ph_cols = ("date", "segment", "brand", "kaznisa", "rrts", "mrc", "opt", "is_current")
        self._price_tree = ttk.Treeview(ph_wrap, columns=ph_cols, show="headings",
                                         style="Analytics.Treeview", selectmode="browse",
                                         height=14)
        for col, heading, width, anchor in [
            ("date",       "Дата снимка",    140, "center"),
            ("segment",    "Сегмент",         80, "center"),
            ("brand",      "Бренд",          140, "w"),
            ("kaznisa",    "КазНИИСА",       110, "center"),
            ("rrts",       "РРЦС",           110, "center"),
            ("mrc",        "МРЦ",            110, "center"),
            ("opt",        "Оптовая",        110, "center"),
            ("is_current", "Статус",         100, "center"),
        ]:
            self._price_tree.heading(col, text=heading)
            self._price_tree.column(col, width=width, anchor=anchor, minwidth=50)

        self._price_tree.tag_configure("current",  background="#D4EDDA", foreground="#155724")
        self._price_tree.tag_configure("history",  background="#F8FAFB")
        self._price_tree.tag_configure("history2", background=BG_CARD)

        ph_vsb = ttk.Scrollbar(ph_wrap, orient="vertical",
                                command=self._price_tree.yview)
        self._price_tree.configure(yscrollcommand=ph_vsb.set)
        self._price_tree.grid(row=0, column=0, sticky="nsew")
        ph_vsb.grid(row=0, column=1, sticky="ns")

    # ── Вкладка 5: Аномалии ───────────────────────────────────────────────────

    def _build_anomalies_tab(self, tab):
        tab.grid_columnconfigure(0, weight=1)
        tab.grid_rowconfigure(1, weight=1)

        ctk.CTkLabel(
            tab, font=FONT_SMALL, text_color=TEXT_SECONDARY,
            text=("Изменения цен > 35% при импортах за период. "
                  "Данные появляются после первого импорта с этой версией приложения."),
        ).grid(row=0, column=0, sticky="w", padx=PAD_MD, pady=(12, 0))

        an_wrap = ctk.CTkFrame(tab, fg_color=BG_CARD,
                               corner_radius=RADIUS_MD, border_width=1,
                               border_color="#DDE2E8")
        an_wrap.grid(row=1, column=0, sticky="nsew", padx=PAD_MD, pady=(8, PAD_MD))
        an_wrap.grid_columnconfigure(0, weight=1)
        an_wrap.grid_rowconfigure(0, weight=1)

        an_cols = ("date", "segment", "article", "brand", "field",
                   "old_price", "new_price", "pct_change")
        self._anomaly_tree = ttk.Treeview(an_wrap, columns=an_cols, show="headings",
                                           style="Analytics.Treeview", selectmode="browse",
                                           height=16)
        for col, heading, width, anchor in [
            ("date",       "Дата",         130, "center"),
            ("segment",    "Сегмент",       80, "center"),
            ("article",    "Артикул",      160, "w"),
            ("brand",      "Бренд",        140, "w"),
            ("field",      "Поле цены",    100, "center"),
            ("old_price",  "Старая цена",  110, "center"),
            ("new_price",  "Новая цена",   110, "center"),
            ("pct_change", "Δ%",            90, "center"),
        ]:
            self._anomaly_tree.heading(col, text=heading)
            self._anomaly_tree.column(col, width=width, anchor=anchor, minwidth=50)

        self._anomaly_tree.tag_configure("up",   background="#F8D7DA", foreground="#721C24")
        self._anomaly_tree.tag_configure("down", background="#D4EDDA", foreground="#155724")
        self._anomaly_tree.tag_configure("odd",  background="#F8FAFB")
        self._anomaly_tree.tag_configure("even", background=BG_CARD)

        an_vsb = ttk.Scrollbar(an_wrap, orient="vertical",
                                command=self._anomaly_tree.yview)
        self._anomaly_tree.configure(yscrollcommand=an_vsb.set)
        self._anomaly_tree.grid(row=0, column=0, sticky="nsew")
        an_vsb.grid(row=0, column=1, sticky="ns")

    # ── Загрузка данных ───────────────────────────────────────────────────────

    def _on_period_change(self, label: str):
        for lbl, days in PERIOD_OPTIONS:
            if lbl == label:
                self._period = days
                self._custom_days_var.set("")
                break
        self.load()

    def _apply_custom_period(self):
        try:
            days = int(self._custom_days_var.get().strip())
            if days < 1:
                raise ValueError
            self._period = days
            self._period_seg.set("")   # снять выбор стандартных кнопок
            self.load()
        except ValueError:
            pass

    def load(self):
        """Загружает все блоки аналитики (кроме поиска цен)."""
        if self._loading:
            return
        self._loading = True
        self._status_lbl.configure(text=f"⏳  Загрузка аналитики за {self._period} дн.…")
        for card in self._cards:
            card.set("…")
        self._avg_lbl.configure(text="")
        self._tree.delete(*self._tree.get_children())
        self._empty_lbl.grid_remove()
        self._brand_tree.delete(*self._brand_tree.get_children())
        self._brand_status_lbl.configure(text="")
        for lbl in self._ai_val_lbls:
            lbl.configure(text="…")
        self._ai_tree.delete(*self._ai_tree.get_children())
        self._anomaly_tree.delete(*self._anomaly_tree.get_children())
        threading.Thread(target=self._fetch, daemon=True).start()

    def _fetch(self):
        try:
            summary   = self.api.get_analytics_summary(self._period)
            kpi       = self.api.get_analytics_kpi(self._period)
            brands    = self.api.get_analytics_brands(self._period)
            ai_eff    = self.api.get_analytics_ai_efficiency(self._period)
            anomalies = self.api.get_analytics_anomalies(self._period)
            self.after(0, lambda: self._render(summary, kpi, brands, ai_eff, anomalies))
        except Exception as e:
            msg = str(e)
            self.after(0, lambda: self._on_error(msg))
        finally:
            self._loading = False

    # ── Рендер ───────────────────────────────────────────────────────────────

    def _render(self, summary: dict, kpi: dict, brands: dict,
                ai_eff: dict, anomalies: dict):
        self._last_summary   = summary
        self._last_kpi       = kpi
        self._last_brands    = brands
        self._last_ai_eff    = ai_eff
        self._last_anomalies = anomalies
        self._status_lbl.configure(text="")

        # Карточки обзора
        for card, v in zip(self._cards, [
            summary.get("total_kp",       0),
            summary.get("total_items",    0),
            summary.get("active_managers", 0),
            summary.get("total_errors",   0),
        ]):
            card.set(f"{v:,}".replace(",", " "))
        avg = summary.get("avg_items_per_kp", 0)
        self._avg_lbl.configure(text=f"Среднее позиций на КП: {avg}")

        # KPI таблица
        self._tree.delete(*self._tree.get_children())
        managers = kpi.get("managers", [])
        if not managers:
            self._empty_lbl.grid(row=4, column=0, pady=8)
        else:
            self._empty_lbl.grid_remove()
            for i, m in enumerate(managers):
                self._tree.insert("", "end",
                                  tags=("odd" if i % 2 else "even",), values=(
                    m.get("full_name",     "—"),
                    m.get("total_kp",       0),
                    m.get("total_items",    0),
                    m.get("avg_items",      0),
                    m.get("last_activity", "—"),
                ))

        self._render_brands(brands)
        self._render_ai(ai_eff)
        self._render_anomalies(anomalies)

        # Обновить заголовок вкладки аномалий с бейджем
        cnt = anomalies.get("count", 0)
        badge = f" ({cnt})" if cnt else ""
        try:
            self._tabs.rename(self._T_ANOMALIES, f"⚠  Аномалии{badge}")
        except Exception:
            pass

    def _render_brands(self, data: dict):
        self._brand_tree.delete(*self._brand_tree.get_children())
        for i, b in enumerate(data.get("top_brands", [])):
            self._brand_tree.insert("", "end",
                                    tags=("odd" if i % 2 else "even",), values=(
                b.get("brand", "—"),
                f"{b.get('corrections', 0):,}".replace(",", " "),
                f"{b.get('pct', 0.0):.1f}%",
                b.get("unique_users", 0),
            ))
        by_status = data.get("by_original_status", {})
        total = data.get("total_corrections", 0)
        if by_status:
            parts = [
                f"{_STATUS_NAMES.get(k, k)}: {v:,}".replace(",", " ")
                for k, v in sorted(by_status.items(), key=lambda x: -x[1])
            ]
            self._brand_status_lbl.configure(
                text=f"Исходный статус исправлений (всего {total:,}): ".replace(",", " ")
                     + "  |  ".join(parts))
        else:
            self._brand_status_lbl.configure(
                text="Нет данных об исправлениях за выбранный период")

    def _render_ai(self, data: dict):
        vals = [
            f"{data.get('total_items', 0):,}".replace(",", " "),
            f"{data.get('total_corrections', 0):,}".replace(",", " "),
            f"{data.get('correction_rate', 0.0):.1f}%",
            f"{data.get('indexing_rate', 0.0):.1f}%",
        ]
        for lbl, v in zip(self._ai_val_lbls, vals):
            lbl.configure(text=v)
        self._ai_tree.delete(*self._ai_tree.get_children())
        by_status = data.get("by_status", [])
        total = sum(s.get("count", 0) for s in by_status)
        for i, s in enumerate(by_status):
            cnt = s.get("count", 0)
            pct = round(cnt / total * 100, 1) if total else 0.0
            self._ai_tree.insert("", "end",
                                  tags=("odd" if i % 2 else "even",), values=(
                _STATUS_NAMES.get(s.get("status", "?"), s.get("status", "?")),
                f"{cnt:,}".replace(",", " "),
                f"{pct:.1f}%",
            ))

    def _render_anomalies(self, data: dict):
        self._anomaly_tree.delete(*self._anomaly_tree.get_children())
        for a in data.get("anomalies", []):
            pct = a.get("pct_change", 0)
            tag = "up" if pct > 0 else "down"
            field_name = _PRICE_FIELDS.get(a.get("field", ""), a.get("field", "—"))
            pct_str = f"{a.get('direction', '')}{abs(pct):.1f}%"
            self._anomaly_tree.insert("", "end", tags=(tag,), values=(
                a.get("date",      "—"),
                a.get("segment",   "—").upper(),
                a.get("article",   "—"),
                a.get("brand",     "—"),
                field_name,
                f"{a.get('old_price', 0):,.2f}".replace(",", " "),
                f"{a.get('new_price', 0):,.2f}".replace(",", " "),
                pct_str,
            ))

    # ── Поиск истории цен ────────────────────────────────────────────────────

    def _search_price_history(self):
        article = self._price_article_var.get().strip()
        if not article:
            return
        segment = self._price_segment_var.get()
        if segment == "Все":
            segment = None
        self._price_status_lbl.configure(text=f"⏳  Ищу «{article}»…")
        self._price_tree.delete(*self._price_tree.get_children())
        threading.Thread(target=self._fetch_price_history,
                         args=(article, segment), daemon=True).start()

    def _fetch_price_history(self, article: str, segment):
        try:
            data = self.api.get_price_history(article, segment)
            self.after(0, lambda: self._render_price_history(data))
        except Exception as e:
            msg = str(e)
            self.after(0, lambda: self._price_status_lbl.configure(
                text=f"⚠️  Ошибка: {msg}"))

    def _render_price_history(self, data: dict):
        self._price_tree.delete(*self._price_tree.get_children())
        records = data.get("records", [])
        if not records:
            self._price_status_lbl.configure(
                text=f"Нет данных для «{data.get('article', '')}»")
            return

        brand = records[-1].get("brand", "—")
        name  = records[-1].get("name",  "—")[:80]
        self._price_status_lbl.configure(
            text=f"{brand} — {name}  |  Записей: {len(records)}"
                 f"  (архивных: {sum(1 for r in records if not r.get('is_current'))},"
                 f" текущих: {sum(1 for r in records if r.get('is_current'))})")

        hist_idx = 0
        for r in records:
            is_curr = r.get("is_current", False)
            tag = "current" if is_curr else ("history" if hist_idx % 2 == 0 else "history2")
            if not is_curr:
                hist_idx += 1

            def fmt_price(v):
                return f"{v:,.2f}".replace(",", " ") if v else "—"

            self._price_tree.insert("", "end", tags=(tag,), values=(
                r.get("date",    "—"),
                r.get("segment", "—").upper(),
                r.get("brand",   "—"),
                fmt_price(r.get("kaznisa")),
                fmt_price(r.get("rrts")),
                fmt_price(r.get("mrc")),
                fmt_price(r.get("opt")),
                "✅ Текущая" if is_curr else "📁 Архив",
            ))

    # ── Экспорт в Excel (#6) ──────────────────────────────────────────────────

    def _export_to_excel(self):
        if not _HAS_OPENPYXL:
            messagebox.showerror("Экспорт", "openpyxl не установлен. "
                                 "Установите: pip install openpyxl")
            return
        if not self._last_summary:
            messagebox.showinfo("Экспорт", "Сначала загрузите данные аналитики.")
            return

        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel файл", "*.xlsx")],
            initialfile=f"analytics_{datetime.date.today()}.xlsx",
            title="Сохранить аналитику",
        )
        if not path:
            return

        try:
            self._write_excel(path)
            messagebox.showinfo("Экспорт", f"Файл сохранён:\n{path}")
        except Exception as e:
            messagebox.showerror("Ошибка экспорта", str(e))

    def _write_excel(self, path: str):
        wb = openpyxl.Workbook()

        HDR_FILL   = PatternFill("solid", fgColor="1E3D6B")
        HDR_FONT   = Font(color="FFFFFF", bold=True)
        HDR_ALIGN  = Alignment(horizontal="center", vertical="center")
        CURR_FILL  = PatternFill("solid", fgColor="D4EDDA")
        ODD_FILL   = PatternFill("solid", fgColor="F8FAFB")

        def _sheet(name: str, headers: list[str], rows: list[list]) -> None:
            ws = wb.create_sheet(name)
            ws.append(headers)
            for cell in ws[1]:
                cell.fill  = HDR_FILL
                cell.font  = HDR_FONT
                cell.alignment = HDR_ALIGN
            for i, row in enumerate(rows):
                ws.append(row)
                if i % 2 == 0:
                    for cell in ws[i + 2]:
                        cell.fill = ODD_FILL
            for col in ws.columns:
                max_len = max((len(str(c.value or "")) for c in col), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

        period = self._period
        s = self._last_summary

        # Лист 1 — Обзор
        ws_ov = wb.active
        ws_ov.title = "Обзор"
        ws_ov["A1"] = f"Аналитика за {period} дней"
        ws_ov["A1"].font = Font(bold=True, size=14)
        ws_ov["A2"] = f"Дата выгрузки: {datetime.date.today()}"
        for cell in ws_ov[4]:
            cell.fill = HDR_FILL
            cell.font = HDR_FONT
        ws_ov.append([])
        ws_ov.append(["Показатель", "Значение"])
        for cell in ws_ov[ws_ov.max_row]:
            cell.fill = HDR_FILL
            cell.font = HDR_FONT
        for label, key in [
            ("КП за период",      "total_kp"),
            ("Позиций обработано", "total_items"),
            ("Активных менеджеров", "active_managers"),
            ("Ошибок",            "total_errors"),
            ("Ср. позиций на КП", "avg_items_per_kp"),
        ]:
            ws_ov.append([label, s.get(key, 0)])

        # Лист 2 — Менеджеры
        _sheet("Менеджеры",
               ["Менеджер", "КП", "Позиций", "Ср. позиций/КП", "Последняя активность"],
               [[m.get("full_name", "—"),
                 m.get("total_kp", 0),
                 m.get("total_items", 0),
                 m.get("avg_items", 0),
                 m.get("last_activity", "—")]
                for m in self._last_kpi.get("managers", [])])

        # Лист 3 — Бренды
        _sheet("Бренды",
               ["Бренд", "Исправлений", "%", "Пользователей"],
               [[b.get("brand", "—"),
                 b.get("corrections", 0),
                 b.get("pct", 0),
                 b.get("unique_users", 0)]
                for b in self._last_brands.get("top_brands", [])])

        # Лист 4 — ИИ-подбор
        _sheet("ИИ-подбор",
               ["Исходный статус", "Исправлений", "%"],
               [[_STATUS_NAMES.get(s.get("status", "?"), s.get("status", "?")),
                 s.get("count", 0),
                 round(s.get("count", 0) /
                       max(sum(x.get("count", 0) for x in
                               self._last_ai_eff.get("by_status", [x])), 1) * 100, 1)]
                for s in self._last_ai_eff.get("by_status", [])])

        # Лист 5 — Аномалии
        _sheet("Аномалии цен",
               ["Дата", "Сегмент", "Артикул", "Бренд", "Поле",
                "Старая цена", "Новая цена", "Δ%"],
               [[a.get("date", "—"), a.get("segment", "—").upper(),
                 a.get("article", "—"), a.get("brand", "—"),
                 _PRICE_FIELDS.get(a.get("field", ""), a.get("field", "—")),
                 a.get("old_price", 0), a.get("new_price", 0),
                 f"{a.get('direction', '')}{abs(a.get('pct_change', 0)):.1f}%"]
                for a in self._last_anomalies.get("anomalies", [])])

        wb.save(path)

    # ── Вспомогательные ───────────────────────────────────────────────────────

    def _on_error(self, msg: str):
        self._status_lbl.configure(text=f"⚠️  Ошибка загрузки: {msg}")
        for card in self._cards:
            card.set("—")

    def _sort_tree(self, col: str, reverse: bool):
        data = [(self._tree.set(k, col), k) for k in self._tree.get_children("")]
        try:
            data.sort(key=lambda x: float(x[0].replace(" ", "")), reverse=reverse)
        except ValueError:
            data.sort(key=lambda x: x[0].lower(), reverse=reverse)
        for idx, (_, k) in enumerate(data):
            self._tree.move(k, "", idx)
            self._tree.item(k, tags=("odd" if idx % 2 else "even",))
        self._tree.heading(col, command=lambda: self._sort_tree(col, not reverse))

    # ── Вспомогательное: вставка в CTkEntry ──────────────────────────────────

    @staticmethod
    def _bind_paste(entry: ctk.CTkEntry):
        """Добавляет правую кнопку (Копировать/Вставить/Вырезать/Всё)
        и явный Ctrl+V для CTkEntry, у которых нет системного контекстного меню."""
        inner = getattr(entry, "_entry", entry)  # tk.Entry внутри CTkEntry

        def _paste(e=None):
            try:
                txt = entry.clipboard_get()
            except Exception:
                return
            try:
                if inner.selection_present():
                    inner.delete("sel.first", "sel.last")
            except Exception:
                pass
            inner.insert("insert", txt)

        def _popup(e):
            menu = tk.Menu(inner, tearoff=0)
            menu.add_command(label="Вырезать",  command=lambda: inner.event_generate("<<Cut>>"))
            menu.add_command(label="Копировать", command=lambda: inner.event_generate("<<Copy>>"))
            menu.add_command(label="Вставить",  command=_paste)
            menu.add_separator()
            menu.add_command(label="Выделить всё",
                             command=lambda: inner.select_range(0, "end"))
            try:
                menu.tk_popup(e.x_root, e.y_root)
            finally:
                menu.grab_release()

        inner.bind("<Control-v>", _paste)
        inner.bind("<Control-V>", _paste)
        inner.bind("<Button-3>",  _popup)

    def after_login(self):
        """Вызывается из main_window после успешного логина."""
        self.after(300, self.load)

"""
Отчёт сверки эксель-базы сегмента с прейскурантом КазНИИСА.

Три листа, по одному на каждый исход сверки — смешивать их на одном листе
бессмысленно, у них разный набор колонок и разные действия по итогу:

    «Итоги»          — сводка по сверке;
    «Обновление цен» — что нашлось по коду или артикулу и какая цена встанет;
    «В общую базу»   — коды прейскуранта без пары в базе сегмента.

Позиции базы, которых прейскурант не знает, отдельным листом не выводятся:
на настоящем прейскуранте туда попадает почти вся база, и лист только
разрастается. Их количество видно в итогах.

Цвета на листе обновления показывают направление: красным — цена выросла,
зелёным — упала. Строки без старой цены выделены отдельно: там не изменение,
а первое заполнение.
"""
import logging
import os
from datetime import datetime
from typing import Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

logger = logging.getLogger(__name__)

FONT_NAME = "Arial"

HDR_FILL  = PatternFill("solid", fgColor="1F3864")
HDR_FONT  = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
BASE_FONT = Font(name=FONT_NAME, size=10)
BOLD_FONT = Font(name=FONT_NAME, size=10, bold=True)
TITLE_FONT = Font(name=FONT_NAME, size=13, bold=True, color="1F3864")

UP_FILL   = PatternFill("solid", fgColor="F8CBAD")   # цена выросла
DOWN_FILL = PatternFill("solid", fgColor="D5E8D4")   # цена упала
NEW_FILL  = PatternFill("solid", fgColor="DEEBF7")   # цены не было
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")

_THIN  = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

MONEY = "#,##0.00"
PCT   = '+0.0%;-0.0%;"—"'


def _header(ws, headers, widths):
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = (f"A1:{openpyxl.utils.get_column_letter(len(headers))}1")


def _put(ws, r, c, value, fmt=None, fill=None, font=None):
    cell = ws.cell(row=r, column=c, value=value)
    cell.font = font or BASE_FONT
    cell.border = BORDER
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


# ─── Листы ───────────────────────────────────────────────────────────────────

def _sheet_changes(wb, rows: List[Dict], applied: bool):
    ws = wb.active
    ws.title = "Обновление цен"
    _header(ws,
            ["Стр.", "Код АГСК", "Артикул", "Наименование (база)",
             "Наименование (прейскурант)", "Бренд", "Ед.",
             "Цена была", "Сметная цена", "НДС", "Цена станет",
             "Разница, ₸", "Разница, %", "Совпало по"],
            [8, 18, 20, 42, 42, 14, 8, 16, 16, 8, 16, 14, 12, 14])

    for i, m in enumerate(rows, start=2):
        old = m.get("price_old")
        if old is None:
            fill = NEW_FILL
        elif m["price_new"] > old:
            fill = UP_FILL
        elif m["price_new"] < old:
            fill = DOWN_FILL
        else:
            fill = None

        _put(ws, i, 1,  m["row"])
        _put(ws, i, 2,  m.get("code_pl") or m.get("code") or "")
        _put(ws, i, 3,  m.get("article", ""))
        _put(ws, i, 4,  m.get("name", ""))
        _put(ws, i, 5,  m.get("name_pl", ""))
        _put(ws, i, 6,  m.get("brand", ""))
        _put(ws, i, 7,  m.get("unit", ""))
        _put(ws, i, 8,  old, MONEY)
        _put(ws, i, 9,  m["price_pl"], MONEY)
        _put(ws, i, 10, 1.16)
        _put(ws, i, 11, m["price_new"], MONEY, fill, BOLD_FONT)
        _put(ws, i, 12, m.get("diff_abs"), MONEY, fill)
        d = m.get("diff_pct")
        _put(ws, i, 13, (d / 100.0) if d is not None else None, PCT, fill)
        _put(ws, i, 14, m.get("how", ""))

    note = ("Цены записаны в исходный файл базы."
            if applied else
            "Режим сравнения: файл базы НЕ изменялся.")
    r = len(rows) + 3
    ws.cell(row=r, column=1, value=note).font = BOLD_FONT
    ws.cell(row=r + 1, column=1,
            value="Красный — цена выросла, зелёный — упала, "
                  "голубой — цены в базе не было.").font = BASE_FONT


def _sheet_to_general(wb, rows: List[Dict], loaded: bool):
    ws = wb.create_sheet("В общую базу")
    _header(ws,
            ["Код АГСК", "Наименование", "Ед.", "Сметная цена",
             "Цена с НДС", "Стр. прейскуранта"],
            [20, 60, 12, 18, 18, 18])
    for i, e in enumerate(rows, start=2):
        price = float(e.get("price") or 0)
        _put(ws, i, 1, e.get("code", ""))
        _put(ws, i, 2, e.get("name", ""))
        _put(ws, i, 3, e.get("unit", ""))
        _put(ws, i, 4, price or None, MONEY)
        _put(ws, i, 5, round(price * 1.16, 2) or None, MONEY, font=BOLD_FONT)
        _put(ws, i, 6, e.get("page"))

    r = len(rows) + 3
    ws.cell(row=r, column=1,
            value=("Позиции загружены в общую базу (сегмент «Общая база (АГСК)»)."
                   if loaded else
                   "Режим сравнения: в базу ничего не загружалось.")).font = BOLD_FONT


def _sheet_summary(wb, stats: Dict, base_path: str, pl_name: str,
                   applied: bool, backup: str):
    ws = wb.create_sheet("Итоги", 0)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 34

    ws["A1"] = "Сверка эксель-базы с прейскурантом КазНИИСА"
    ws["A1"].font = TITLE_FONT
    r = 3

    def line(label, value, bold=False):
        nonlocal r
        ws.cell(row=r, column=1, value=label).font = BOLD_FONT if bold else BASE_FONT
        c = ws.cell(row=r, column=2, value=value)
        c.font = BOLD_FONT if bold else BASE_FONT
        r += 1

    line("Дата сверки", datetime.now().strftime("%d.%m.%Y %H:%M"))
    line("Файл базы", os.path.basename(base_path))
    line("Прейскурант", pl_name)
    line("Режим",
         "Цены применены к файлу"
         if applied else
         "ТОЛЬКО СРАВНЕНИЕ — файл базы не изменялся", True)
    if backup:
        line("Копия до изменений", os.path.basename(backup))
    r += 1

    line("Строк в базе",             stats.get("base_rows", 0), True)
    line("Кодов в прейскуранте",     stats.get("pricelist_codes", 0), True)
    r += 1
    line("Сопоставлено",             stats.get("matched", 0), True)
    line("    по коду АГСК",         stats.get("by_code", 0))
    line("    по артикулу",          stats.get("by_article", 0))
    line("Из них цена изменилась",   stats.get("changed", 0), True)
    r += 1
    line("Не опознано в базе",       stats.get("unmatched_base", 0))
    line("В общую базу",             stats.get("to_general", 0), True)
    r += 1
    line("НДС", stats.get("vat", 1.16))
    r += 1
    ws.cell(row=r, column=1,
            value="Цена в базе = верхнее число ценовой ячейки прейскуранта "
                  "(сметная цена) × НДС.").font = BASE_FONT
    ws.cell(row=r + 1, column=1,
            value="Сопоставление только по коду АГСК и артикулу. "
                  "По наименованию цены не подставляются.").font = BASE_FONT


# ─── Точка входа ─────────────────────────────────────────────────────────────

def build_sync_report(out_path: str, result: Dict, base_path: str,
                      pricelist_name: str, applied: bool,
                      backup: str = "", loaded: bool = False) -> str:
    """Собирает отчёт сверки и сохраняет в out_path. Возвращает путь."""
    wb = openpyxl.Workbook()
    _sheet_changes(wb, result.get("matched", []), applied)
    _sheet_to_general(wb, result.get("unmatched_pricelist", []), loaded)
    _sheet_summary(wb, result.get("stats", {}), base_path, pricelist_name,
                   applied, backup)
    wb.active = 0
    wb.save(out_path)
    logger.info("sync report saved: %s", out_path)
    return out_path


def default_report_path(base_path: str) -> str:
    """Путь отчёта рядом с файлом базы, с датой в имени."""
    folder, fname = os.path.split(base_path)
    stem = os.path.splitext(fname)[0]
    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
    return os.path.join(folder, f"Сверка с прейскурантом — {stem} — {stamp}.xlsx")

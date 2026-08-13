"""
Отчёт сверки цен базы со сметными ценами прейскуранта АГСК.

Лист «Сверка» — построчное сравнение, отсортировано по убыванию отклонения.
Лист «Итоги» — сводка и пояснения к методике.

Цвета:
    красный — отклонение больше порога;
    жёлтый  — единицы измерения различаются (цена в базе за упаковку/бухту,
              в прейскуранте за метр), отклонение у таких строк мнимое.
"""
import logging
from typing import Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

FONT_NAME = "Arial"

HDR_FILL  = PatternFill("solid", fgColor="1F3864")
HDR_FONT  = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
BASE_FONT = Font(name=FONT_NAME, size=10)
RED_FILL  = PatternFill("solid", fgColor="F8CBAD")
RED_FONT  = Font(name=FONT_NAME, size=10, bold=True, color="9C0006")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")

_THIN  = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

MONEY = "#,##0.00"
PCT   = '0.0%;-0.0%;"—"'

HEADERS = [
    "Код АГСК", "Артикул", "Наименование (база)", "Наименование (прейскурант)",
    "Бренд", "Сегмент", "Ед. (база)", "Ед. (прейскурант)",
    "Цена КазНИИСА (база)", "Сметная цена (прейскурант)",
    "Разница, тенге", "Разница, %", "Отклонение", "Единицы совпадают",
    "Стр. прейскуранта",
]
WIDTHS = [20, 18, 44, 44, 14, 10, 12, 16, 20, 24, 16, 12, 14, 18, 16]

SEG_NAMES = {"ss": "Слаботочные", "os": "Освещение", "sil": "Силовые"}


def build_report(path: str, rows: List[Dict], stats: Dict) -> str:
    """Сохраняет отчёт сверки в файл path. Возвращает путь."""
    threshold = float(stats.get("threshold_pct", 5.0))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сверка"

    for c, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"

    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1,  value=r["code"])
        ws.cell(row=i, column=2,  value=r.get("article", ""))
        ws.cell(row=i, column=3,  value=r.get("name_db", ""))
        ws.cell(row=i, column=4,  value=r.get("name_pl", ""))
        ws.cell(row=i, column=5,  value=r.get("brand", ""))
        ws.cell(row=i, column=6,  value=SEG_NAMES.get(r.get("segment", ""),
                                                      r.get("segment", "")))
        ws.cell(row=i, column=7,  value=r.get("unit_db", ""))
        ws.cell(row=i, column=8,  value=r.get("unit_pl", ""))
        ws.cell(row=i, column=9,  value=r["price_db"]).number_format = MONEY
        ws.cell(row=i, column=10, value=r["price_pl"]).number_format = MONEY
        # Разница — формулами, чтобы пересчитывалась при правке цен
        ws.cell(row=i, column=11, value=f"=I{i}-J{i}").number_format = MONEY
        ws.cell(row=i, column=12,
                value=f'=IF(J{i}=0,"",(I{i}-J{i})/J{i})').number_format = PCT
        ws.cell(row=i, column=13,
                value=(f'=IF(L{i}="","",IF(ABS(L{i})>{threshold}/100,'
                       f'"более {threshold:g}%","в норме"))'))
        ws.cell(row=i, column=14, value="да" if r.get("same_unit") else "НЕТ")
        ws.cell(row=i, column=15, value=r.get("page"))

        over = bool(r.get("over"))
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=i, column=c)
            cell.border = BORDER
            cell.font = RED_FONT if over else BASE_FONT
            if over:
                cell.fill = RED_FILL
            elif not r.get("same_unit"):
                cell.fill = WARN_FILL
        ws.cell(row=i, column=1).alignment = Alignment(horizontal="left")

    last = len(rows) + 1
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last}"

    # ─── Итоги ───────────────────────────────────────────────────────────────
    s = wb.create_sheet("Итоги")
    s.column_dimensions["A"].width = 56
    s.column_dimensions["B"].width = 20

    def put(row, label, value, fmt=None, bold=False):
        a = s.cell(row=row, column=1, value=label)
        b = s.cell(row=row, column=2, value=value)
        a.font = Font(name=FONT_NAME, size=10, bold=bold)
        b.font = Font(name=FONT_NAME, size=10, bold=bold)
        if fmt:
            b.number_format = fmt

    s["A1"] = "Сверка цен базы со сметными ценами прейскуранта АГСК"
    s["A1"].font = Font(name=FONT_NAME, size=12, bold=True, color="1F3864")

    segs = stats.get("segments") or []
    put(3,  "Файл прейскуранта",                     stats.get("filename", ""))
    put(4,  "Сегменты базы",
        ", ".join(SEG_NAMES.get(x, x) for x in segs) or "—")
    put(5,  "Порог отклонения",                      threshold / 100, "0.0%")

    put(7,  "Кодов в прейскуранте",                  stats.get("pricelist_codes", 0))
    put(8,  "Товаров в выбранных сегментах",         stats.get("products_total", 0))
    put(9,  "   из них без кода АГСК",               stats.get("without_code", 0))
    put(10, "   код есть, но в прейскуранте не найден", stats.get("not_in_pricelist", 0))
    put(11, "Сопоставлено и сравнено",               stats.get("matched", 0), bold=True)

    put(13, f"Отклонение больше {threshold:g}%",
        f'=COUNTIF(Сверка!M2:M{last},"более {threshold:g}%")' if rows else 0, bold=True)
    put(14, "   цена базы выше прейскуранта",        stats.get("over_higher", 0))
    put(15, "   цена базы ниже прейскуранта",        stats.get("over_lower", 0))
    put(16, "Отклонение в пределах порога",
        f'=COUNTIF(Сверка!M2:M{last},"в норме")' if rows else 0)
    put(18, "Из них с РАЗНЫМИ единицами измерения",  stats.get("unit_mismatch", 0))
    put(19, "   (отклонение мнимое, сравнивать нельзя)", "")

    if rows:
        put(21, "Среднее отклонение, %",  f"=AVERAGE(Сверка!L2:L{last})", PCT)
        put(22, "Максимальное отклонение", f"=MAX(Сверка!L2:L{last})", PCT)
        put(23, "Минимальное отклонение",  f"=MIN(Сверка!L2:L{last})", PCT)

    s["A25"] = "Как читать отчёт"
    s["A25"].font = Font(name=FONT_NAME, size=11, bold=True, color="1F3864")
    notes = [
        f"Красная заливка — отклонение цены больше {threshold:g}%.",
        "Жёлтая заливка — единицы измерения в базе и прейскуранте различаются:",
        "    например, в базе цена за бухту 200 м, а в прейскуранте за метр.",
        "    Отклонение у таких строк мнимое, сравнивать их напрямую нельзя.",
        "Разница, % = (цена базы − сметная цена) / сметная цена.",
        "    Плюс означает, что наша цена выше прейскуранта.",
        "Сметная цена — верхнее число в ячейке прейскуранта;",
        "    нижнее (отпускная цена) в сверке не используется.",
        "Сопоставление строго по коду АГСК, точное совпадение.",
    ]
    for k, line in enumerate(notes, start=26):
        c = s.cell(row=k, column=1, value=line)
        c.font = Font(name=FONT_NAME, size=10,
                      bold=line.endswith(":") and not line.startswith("    "))

    wb.save(path)
    logger.info("pricelist_report: сохранён %s (%d строк)", path, len(rows))
    return path

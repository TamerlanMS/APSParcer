"""
Разбор спецификации в формате Excel (.xlsx / .xlsm).

Ожидаемая структура листа (колонки определяются по строке заголовка):
    Поз. | Наименование и техническая характеристика | Тип, марка, обозначение |
    Код продукции | Поставщик | Ед. измерения | Количество | Масса | Примечание

Особенности формата:
  • Строки-разделы — заполнено только «Наименование» («1.Щитовое оборудование»).
  • Составные позиции — строка с «Поз.» и количеством («компл.»), под ней
    строки-компоненты без номера позиции и без количества.
  • Подбор выполняется по ВСЕМ строкам: и по составной позиции, и по компонентам.

Для обратной записи в исходный файл у каждой позиции сохраняется
`row` — номер строки на листе и `sheet` — имя листа.
"""
import logging
import re
from typing import Dict, List, Optional, Tuple

import openpyxl

logger = logging.getLogger(__name__)


# ─── Распознавание колонок по заголовку ──────────────────────────────────────

_COL_PATTERNS: List[Tuple[str, Tuple[str, ...]]] = [
    ("pos",     ("поз",)),
    ("name",    ("наименование", "техническая характеристика", "наимен")),
    ("article", ("тип", "марка", "обозначение")),
    ("code",    ("код продукции", "код")),
    ("brand",   ("поставщик", "изготовитель", "производитель", "бренд")),
    ("unit",    ("ед. изме", "ед.изм", "единица", "ед изм")),
    ("qty",     ("кол-во", "количество", "коли-чест")),
    ("mass",    ("масса",)),
    ("note",    ("примечание", "приме-чание")),
]

_MAX_HEADER_SCAN = 30      # в скольких первых строках ищем шапку
_MIN_HEADER_HITS = 3       # сколько колонок должно совпасть, чтобы счесть строкой шапки

# Слова-категории: строка с таким текстом и без артикула — раздел, а не позиция.
# Тот же список, что и в pdf_parser.SKIP_KEYWORDS.
SKIP_KEYWORDS = [
    "оборудование",
    "кабели",
    "провода",
    "монтажные",
    "кабеленесущие",
    "материалы и изделия",
    "комплектующ",
]

# Служебные метки в начале строки. Снимаются только эти — произвольный текст
# до двоеточия трогать нельзя: «Щит ЩО-0.1: модульный навесной» — не метка.
_LABEL_WORDS = [
    "на вводе",
    "на линиях",
    "на отходящих линиях",
    "в том числе",
    "в т.ч",
    "в т. ч",
    "дополнительно",
    "состав",
    "в составе",
    "комплект",
    "в комплекте",
    "включая",
    "поставка",
    "устанавливается",
    "примечание",
]
_LABEL_PREFIX_RE = re.compile(
    r"^\s*(?:" + "|".join(re.escape(w) for w in _LABEL_WORDS) + r")[^:]{0,20}:\s*",
    re.IGNORECASE,
)

# Ведущие дефисы, тире и маркеры списка
_LEAD_BULLET_RE = re.compile(r"^[\s\-\u2013\u2014\u2022\u00b7]+")

# Хвостовое количество: «- 3 шт.», «— 12 шт»
_TAIL_QTY_RE = re.compile(
    r"\s*[-\u2013\u2014]\s*(\d+(?:[.,]\d+)?)\s*(?:шт|компл|к-?т)\.?\s*$",
    re.IGNORECASE,
)


def _clean_component_name(text: str):
    """Убирает служебную метку/маркеры и вынимает количество из хвоста.

    Возвращает (очищенное имя, количество или None).
    Метка снимается только если после неё остаётся осмысленный текст.
    """
    s = (text or "").strip()
    if not s:
        return "", None

    qty = None
    m = _TAIL_QTY_RE.search(s)
    if m:
        try:
            qty = float(m.group(1).replace(",", "."))
        except ValueError:
            qty = None
        s = s[:m.start()].strip()

    # Метку снимаем, только если она действительно метка, а не часть названия
    # (после двоеточия должно остаться не меньше 4 символов).
    lm = _LABEL_PREFIX_RE.match(s)
    if lm and len(s) - lm.end() >= 4:
        s = s[lm.end():]

    s = _LEAD_BULLET_RE.sub("", s).strip()
    return s, qty


def _is_description_row(name: str, article: str, code: str,
                        is_component: bool, has_parent: bool) -> bool:
    """True — строка описывает состав родительской позиции, а не товар.

    Подобрать такую строку невозможно: у неё нет ни артикула, ни кода,
    а имя — фрагмент перечисления («На вводе: - Выключатель ...»).
    """
    if article or code:
        return False          # есть чем искать — это позиция
    if not is_component:
        return False          # у строки свой номер позиции или количество

    nm = (name or "").strip()
    # Настоящий раздел описанием не считается, даже если выше были позиции
    if re.match(r"^\d+\s*[.)]\s*\S", nm):
        return False          # «2.Световое оборудование»
    if any(kw in nm.lower() for kw in SKIP_KEYWORDS):
        return False          # «Кабели и провода», «Монтажные изделия»

    if _LABEL_PREFIX_RE.match(nm):
        return True           # явная служебная метка «На вводе:», «На линиях:»
    return has_parent         # фрагмент перечисления под известной позицией


def _norm(v) -> str:
    """Нормализует значение ячейки в строку без лишних пробелов и переносов."""
    if v is None:
        return ""
    s = str(v).replace("\n", " ").replace("\r", " ")
    return re.sub(r"\s+", " ", s).strip()


def _detect_header(ws, max_col: int) -> Tuple[Optional[int], Dict[str, int]]:
    """Ищет строку заголовка и возвращает (номер строки, {ключ: номер колонки})."""
    best_row, best_map, best_hits = None, {}, 0

    for r in range(1, min(ws.max_row, _MAX_HEADER_SCAN) + 1):
        cells = {c: _norm(ws.cell(row=r, column=c).value).lower()
                 for c in range(1, max_col + 1)}
        if not any(cells.values()):
            continue

        col_map: Dict[str, int] = {}
        for key, patterns in _COL_PATTERNS:
            if key in col_map:
                continue
            for c, text in cells.items():
                if not text or c in col_map.values():
                    continue
                if any(p in text for p in patterns):
                    col_map[key] = c
                    break

        hits = len(col_map)
        if hits > best_hits:
            best_row, best_map, best_hits = r, col_map, hits

    if best_hits < _MIN_HEADER_HITS:
        return None, {}
    return best_row, best_map


def _fallback_map(max_col: int) -> Dict[str, int]:
    """Позиционная раскладка, если шапку распознать не удалось."""
    order = ["pos", "name", "article", "code", "brand", "unit", "qty", "mass", "note"]
    return {k: i + 1 for i, k in enumerate(order) if i + 1 <= max_col}


def _parse_qty(v) -> float:
    """Количество: допускает «12», «12,5», «12.5», мусор → 0."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        try:
            return float(v)
        except (TypeError, ValueError):
            return 0.0
    s = _norm(v).replace(",", ".")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    if not m:
        return 0.0
    try:
        return float(m.group())
    except ValueError:
        return 0.0


def _is_section_heading(pos: str, name: str, article: str,
                        code: str, qty: float, unit: str) -> bool:
    """Строка-раздел: есть только наименование, всё остальное пусто."""
    if not name:
        return False
    if article or code or unit or qty:
        return False
    low = name.lower()
    if any(kw in low for kw in SKIP_KEYWORDS) and not article and not code:
        return True
    if pos:
        # «1.» / «2.» у раздела допустимы, но развёрнутый номер «1.1» — это позиция
        return not re.match(r"^\d+\.\d", pos)
    # Типовой вид раздела: «1.Щитовое оборудование», «2. Световое оборудование»
    return bool(re.match(r"^\d+\s*[.)]\s*\S", name)) or len(name) <= 80


def parse_spec_excel(path: str, sheet_name: Optional[str] = None) -> List[Dict]:
    """
    Разбирает файл спецификации в список позиций.

    Возвращает элементы в формате, который принимает matcher.match_items:
        pos, name_raw, article_raw, kaznisa_code_raw, qty, unit, is_heading
    плюс служебные поля для обратной записи:
        row (номер строки листа), sheet (имя листа), brand_raw, note_raw,
        is_component (строка-компонент составной позиции), parent_pos.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]

    max_col = min(ws.max_column or 1, 30)
    header_row, col_map = _detect_header(ws, max_col)
    if not col_map:
        header_row, col_map = 1, _fallback_map(max_col)
        logger.warning("spec_parser: шапка не распознана в %s, использую позиционную раскладку", path)
    else:
        logger.info("spec_parser: шапка в строке %s, колонки=%s", header_row, col_map)

    def cell(r: int, key: str) -> str:
        c = col_map.get(key)
        return _norm(ws.cell(row=r, column=c).value) if c else ""

    def cell_raw(r: int, key: str):
        c = col_map.get(key)
        return ws.cell(row=r, column=c).value if c else None

    items: List[Dict] = []
    skipped_desc = 0       # сколько описательных строк отброшено
    last_pos = ""          # номер последней «полноценной» позиции — родитель компонентов
    start = (header_row or 1) + 1

    for r in range(start, (ws.max_row or 0) + 1):
        pos     = cell(r, "pos")
        name    = cell(r, "name")
        article = cell(r, "article")
        code    = cell(r, "code")
        brand   = cell(r, "brand")
        unit    = cell(r, "unit")
        note    = cell(r, "note")
        qty     = _parse_qty(cell_raw(r, "qty"))

        # Полностью пустая строка — пропускаем
        if not any((pos, name, article, code, brand, unit, note)) and not qty:
            continue

        # ── Расшифровка состава родительской позиции — в подбор не идёт ─────
        # Проверяем ДО признака раздела: такая строка содержит только имя и
        # иначе была бы принята за заголовок.
        if _is_description_row(name, article, code,
                               not pos and not qty, bool(last_pos)):
            skipped_desc += 1
            continue

        # ── Раздел ──────────────────────────────────────────────────────────
        if _is_section_heading(pos, name, article, code, qty, unit):
            items.append({
                "pos":              pos,
                "name_raw":         name,
                "article_raw":      "",
                "kaznisa_code_raw": "",
                "qty":              0,
                "unit":             "",
                "is_heading":       True,
                "row":              r,
                "sheet":            ws.title,
                "brand_raw":        "",
                "note_raw":         note,
                "is_component":     False,
                "parent_pos":       "",
            })
            continue

        # Строка без чего-либо опознаваемого — не позиция
        if not (article or code or name):
            continue

        is_component = not pos and not qty
        if pos:
            last_pos = pos

        # Служебные метки мешают подбору — чистим имя, как в PDF-парсере
        clean_name, tail_qty = _clean_component_name(name)
        if not clean_name:
            clean_name = name
        if not qty and tail_qty:
            qty = tail_qty

        items.append({
            "pos":              pos or "",
            "name_raw":         clean_name,
            "name_orig":        name,
            "article_raw":      article,
            "kaznisa_code_raw": code,
            "qty":              qty if qty else 1,
            "unit":             unit or "шт.",
            "is_heading":       False,
            "row":              r,
            "sheet":            ws.title,
            "brand_raw":        brand,
            "note_raw":         note,
            "is_component":     is_component,
            "parent_pos":       last_pos if is_component else "",
        })

    real = sum(1 for i in items if not i.get("is_heading"))
    logger.info("spec_parser: %s — всего строк %d, позиций %d, разделов %d, "
                "отброшено описательных %d",
                path, len(items), real, len(items) - real, skipped_desc)
    return items


def get_sheet_names(path: str) -> List[str]:
    """Список листов файла — для выбора нужного, если их несколько."""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()
